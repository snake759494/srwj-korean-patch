# -*- coding: utf-8 -*-
"""삼성효님 번역(초반 1~7화)을 제1 기준으로 시나리오 엑셀에 반영.

md 형식: 한 블록에 `화자「일본어」` 다음 `화자「한국어」`.
매칭: md 일본어 ↔ 엑셀 '일본어(번역TXT)'(col9, 스크립트 원문).
치환: 엑셀 '일본어(ROM 디코딩)'(col11)에는 이름이 변수 ①②③④⑤⑥ 로 들어 있다.
      col9↔col11 을 diff 해 (변수 ← 스크립트이름) 을 얻고,
      md 한국어에서 그 이름의 한국어 표기를 변수로 되돌린다.
      → 제어문자(변수) 규칙을 충실히 보존한다.

사용: python apply_samsunghyo.py [--write]
"""
import sys, re, glob, os, difflib, collections, json

HERE = os.path.dirname(os.path.abspath(__file__))
XLSX = os.path.join(HERE, 'srwj_matched_all_0625.xlsx')
MDDIR = os.path.join(os.path.dirname(HERE), '삼성효님 번역')
WRITE = '--write' in sys.argv

# md 파일 → 대상 archive (매칭 분포로 확인됨)
FILE_ARC = {
    '남주 1화': 1, '남주 2화': 3, '여주 1화': 5, '여주 2화': 7,
    '남주 3화': 9, '여주 3화': 9, '남주 4화': 11, '여주 4화': 11,
    '남주 5화': 13, '여주 5화': 13, '남주 6화 완': 15, '여주 6화 완': 15,
    '남주 7화 완': 17, '여주 7화 완': 17,
}
# 스크립트 이름 → 한국어 표기 (긴 것 먼저 치환)
JP2KO = [
    ('紫雲　統夜', ['시운 토우야', '시운　토우야', '시운토우야']),
    ('紫雲 統夜',       ['시운 토우야', '시운　토우야', '시운토우야']),
    ('カルヴィナ・クーランジュ', ['칼비나 크란쥬', '칼비나·크란쥬', '칼비나・크란쥬', '칼비나크란쥬']),
    ('カルヴィナさん', ['칼비나 씨', '칼비나씨', '칼비나']),
    ('B・ブリガンディ', ['B・브리간디', 'B·브리간디', 'B 브리간디', '브리간디']),
    ('ベルゼルート',   ['벨제루트']),
    ('グランティード', ['그랑티드', '그란티드']),
    ('クーランジュ',   ['크란쥬', '쿠란주', '쿠랑주']),
    ('ク―ランジュ',   ['크란쥬']),
    ('カルヴィナ',     ['칼비나']),
    ('統夜',           ['토우야', '토야']),
    ('紫雲',           ['시운']),
]
CIRC = set('①②③④⑤⑥')
BR = re.compile(r'([^\n「]{0,15})「((?:[^」])*)」', re.S)
# 폰트(KS X 1001 2350자) 밖 글자 = 표시 불가. 원문 오타 교정.
TYPO = [('떄', '때'), ('떈', '땐')]
# 말줄임표 뒤 띄어쓰기. md 원문은 '어…쿄코구나' 처럼 붙여 쓰는데, '…'는 게임에서
# '・・・'(3칸)로 펼쳐지므로 첫 단어가 첫 줄 예산(7칸)을 넘어 단어 중간이 잘린다.
# 뜻은 그대로 두고 띄어쓰기만 넣어 줄바꿈이 어절 경계에서 일어나게 한다.
ELLIPSIS = re.compile(r'(…|\.\.\.)(?=[^\s\.…?!,、。」』\)）])')
# 초반부 archive 집합 (교차 매칭 허용 범위)
EARLY = (1, 3, 5, 7, 9, 11, 13, 15, 17)
# 번역가가 이름 대신 대명사를 써서 변수가 사라진 곳 — 수동 보정
#   ① ⑤ 뒤 조사는 남주/여주·기체 3종 모두 받침이 없어 '를/라면' 이 안전하다.
MANUAL = {
    205:  '해냈어! 굉장해, ③!',
    2822: '그치만 ①라면 할 수 있을 거라고 생각했는걸.',
}
MANUAL_SUB = {273: [('그걸', '⑤를')]}


def norm(s):
    s = s.replace('〜', '～').replace('－', 'ー').replace('･', '・')
    return re.sub(r'\s+', '', s)


def parse_md(path):
    txt = open(path, encoding='utf-8').read()
    txt = re.sub(r'!\[[^\]]*\]\([^)]*\)', '', txt)      # 이미지 제거
    out = []
    for blk in re.split(r'\n\s*\n', txt):
        blk = blk.strip()
        if not blk:
            continue
        ms = BR.findall(blk)
        if len(ms) >= 2:
            out.append((ms[0][1], ms[1][1]))            # (일본어, 한국어)
    return out


def var_pairs(script_jp, rom_jp):
    """col9↔col11 diff 로 (변수, 스크립트이름) 목록.

    dst 가 '③さん' 처럼 변수+조사인 경우도 잡되, 변수 하나만 뽑는다.
    """
    pairs = []
    sm = difflib.SequenceMatcher(None, script_jp, rom_jp, autojunk=False)
    for tag, i1, i2, j1, j2 in sm.get_opcodes():
        if tag == 'replace':
            src, dst = script_jp[i1:i2], rom_jp[j1:j2]
            vs = [c for c in dst if c in CIRC]
            if len(vs) == 1 and len(dst) <= 4 and 1 <= len(src) <= 14:
                pairs.append((vs[0], src))
    return pairs


def rom_ok(script_jp, rom_jp):
    """col9(스크립트)와 col11(ROM)이 같은 대사인지. 이름 변수 치환만 다른 경우 통과.

    둘이 크게 다르면 그 행은 매칭이 잘못된 것이므로 md 번역을 넣으면 안 된다.
    """
    if not script_jp or not rom_jp:
        return False
    a, b = norm(script_jp), norm(rom_jp)
    if a == b:
        return True
    return difflib.SequenceMatcher(None, a, b, autojunk=False).ratio() >= 0.80


def to_vars(ko, pairs, rom_jp):
    """md 한국어에서 이름을 변수로 되돌린다."""
    out = ko
    used = []
    for var, name in pairs:
        cands = None
        for jp, kos in JP2KO:
            if jp == name or (len(name) > 2 and jp in name):
                cands = kos
                break
        if not cands:
            continue
        for k in cands:
            if k in out:
                out = out.replace(k, var, 1)
                used.append((var, k))
                break
    return out, used


def main():
    from openpyxl import load_workbook
    wb = load_workbook(XLSX)
    ws = wb['매칭 결과']
    rows = list(ws.iter_rows(min_row=2, values_only=True))

    # archive별 norm(col9) → 행번호들
    idx = collections.defaultdict(list)
    for i, r in enumerate(rows, 2):
        if r[8]:
            idx[(r[0], norm(str(r[8])))].append(i)

    stats = collections.Counter()
    plan = {}          # 행번호 → (새KO, 출처, 경고)
    conflicts = []
    warn = []

    for path in sorted(glob.glob(os.path.join(MDDIR, '*.md'))):
        base = os.path.splitext(os.path.basename(path))[0]
        arc = FILE_ARC.get(base)
        if arc is None:
            print(f"[스킵] archive 미지정: {base}")
            continue
        pairs = parse_md(path)
        for jp, ko in pairs:
            stats['md_pairs'] += 1
            key = (arc, norm(jp))
            hits = idx.get(key)
            if not hits:
                # 지정 archive 에 없으면 초반부 archive 전체에서 유일하게 걸릴 때만 허용.
                # 짧은 대사('네', '화성?' 등)는 우연 일치가 잦아 제외한다.
                alt = ([h for a in EARLY for h in idx.get((a, norm(jp)), [])]
                       if len(norm(jp)) >= 10 else [])
                if len(alt) == 1:
                    hits = alt
                    stats['cross_arc'] += 1
                else:
                    stats['unmatched'] += 1
                    continue
            if len(hits) > 1:
                stats['ambiguous'] += 1
                continue
            rn = hits[0]
            r = rows[rn - 2]
            rom_jp = str(r[10]) if r[10] else ''
            script_jp = str(r[8]) if r[8] else ''
            if not rom_ok(script_jp, rom_jp):
                stats['rom_mismatch'] += 1     # 스크립트↔ROM 불일치 행 → 넣지 않는다
                continue
            newko = ko.strip()
            for a, b in TYPO:
                newko = newko.replace(a, b)
            newko = ELLIPSIS.sub(r'\1 ', newko)
            if not newko:
                stats['empty_ko'] += 1
                continue
            # 변수 복원
            vp = var_pairs(script_jp, rom_jp)
            if vp:
                newko, used = to_vars(newko, vp, rom_jp)
            # 수동 보정(대명사로 번역돼 변수가 사라진 곳)
            if rn in MANUAL:
                newko = MANUAL[rn]
            for a, b in MANUAL_SUB.get(rn, []):
                newko = newko.replace(a, b, 1)
            # 변수 개수 검증
            for c in CIRC:
                need, have = rom_jp.count(c), newko.count(c)
                if need and not have:
                    warn.append((rn, arc, c, need, have, rom_jp[:40], newko[:40]))
                    stats['var_missing'] += 1
            if rn in plan and plan[rn][0] != newko:
                conflicts.append((rn, plan[rn][1], plan[rn][0], base, newko))
                stats['conflict'] += 1
                continue
            plan[rn] = (newko, base)
            stats['matched'] += 1

    print("=== 통계 ===")
    for k in ('md_pairs', 'matched', 'unmatched', 'ambiguous', 'cross_arc', 'rom_mismatch', 'empty_ko', 'conflict', 'var_missing'):
        print(f"  {k:12}: {stats[k]}")
    print(f"  적용 대상 행수: {len(plan)}")

    # 변경 없는 것 제외
    changed = {rn: v for rn, v in plan.items() if str(rows[rn - 2][9] or '') != v[0]}
    print(f"  실제 변경될 행: {len(changed)}")

    print(f"\n=== 변수 누락 경고 {len(warn)}건 (상위 12) ===")
    for rn, arc, c, need, have, rj, nk in warn[:12]:
        print(f"  행{rn} a{arc} {c} ROM {need}개 → KO {have}개")
        print(f"     ROM:{rj!r}")
        print(f"     KO :{nk!r}")

    print(f"\n=== 남주/여주 충돌 {len(conflicts)}건 (상위 6) ===")
    for rn, f1, k1, f2, k2 in conflicts[:6]:
        print(f"  행{rn}  [{f1}] {k1[:40]!r}")
        print(f"        [{f2}] {k2[:40]!r}")

    print("\n=== 변경 샘플 12건 ===")
    for rn in sorted(changed)[:12]:
        print(f"  행{rn} ({changed[rn][1]})")
        print(f"     전: {str(rows[rn-2][9])[:56]!r}")
        print(f"     후: {changed[rn][0][:56]!r}")

    if WRITE:
        for rn, (ko, src) in changed.items():
            ws.cell(rn, 10).value = ko
        wb.save(XLSX)
        print(f"\n✓ 저장: {len(changed)}행 반영")
    else:
        print("\n(분석 전용 — 반영하려면 --write)")
    json.dump({str(k): v[0] for k, v in changed.items()},
              open(os.path.join(HERE, '_samsunghyo_plan.json'), 'w', encoding='utf-8'),
              ensure_ascii=False, indent=1)


if __name__ == '__main__':
    main()
