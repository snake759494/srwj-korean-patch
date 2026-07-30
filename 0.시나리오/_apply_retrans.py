# -*- coding: utf-8 -*-
"""재번역 결과(out_NNN.json) 수집 → 기계 검증 → 엑셀 반영.

검증 항목(행 단위):
  V1 변수 ①~⑥ 개수가 ROM 원문과 일치
  V2 일본어 문자(가나/한자) 잔존 없음
  V3 한글이 KS X 1001 2350자 내
  V4 비어있지 않음 / 원문 그대로 복사 아님
  V5 화자「 프리픽스 없음
  자동수정: 말줄임표 뒤 띄어쓰기, _x000D_ 제거, 앞뒤 공백
길이 경고(줄수 > 원문+3)는 수집만(fail 아님).

사용: python _apply_retrans.py <retrans_dir> [--write]
"""
import sys, os, json, re, glob, collections

sys.stdout.reconfigure(encoding='utf-8')
HERE = os.path.dirname(os.path.abspath(__file__))
DIR = sys.argv[1]
WRITE = '--write' in sys.argv

KS = {bytes([0xB0 + i // 94, 0xA1 + i % 94]).decode('euc-kr') for i in range(2350)}
CIRC = '①②③④⑤⑥'
JPCH = re.compile(r'[ぁ-ゖァ-ヺ一-鿿ー]')
ELL = re.compile(r'(…|\.\.\.)(?=[^\s\.…?!,、。」』\)）])')

from openpyxl import load_workbook
wb = load_workbook(os.path.join(HERE, 'srwj_matched_all_0625.xlsx'))
ws = wb['매칭 결과']
rows = list(ws.iter_rows(min_row=2, values_only=True))

# 대상 집합 재구성
targets = {}
for i, r in enumerate(rows, 2):
    ko = str(r[9]) if r[9] else ''
    jp = str(r[10]) if r[10] else ''
    human = bool(r[8] and str(r[8]).strip())
    if ko.strip() and jp.strip() and not human:
        targets[i] = jp

# 수집
got = {}
dupdiff = 0
for f in sorted(glob.glob(os.path.join(DIR, 'out_*.json'))):
    try:
        d = json.load(open(f, encoding='utf-8'))
    except Exception as e:
        print(f"[깨진 파일] {os.path.basename(f)}: {e}")
        continue
    for row in d.get('rows', []):
        rn = int(row['rn']); ko = str(row.get('ko', ''))
        if rn in got and got[rn] != ko:
            dupdiff += 1
        got[rn] = ko

print(f"대상 {len(targets)}행 / 수집 {len(got)}행 / 중복상이 {dupdiff}")
missing = sorted(set(targets) - set(got))
extra = sorted(set(got) - set(targets))
print(f"누락 {len(missing)} / 대상외 {len(extra)}")

fails = collections.defaultdict(list)
warn_len = []
ok = {}
from srwj_wrap import fit_turn_lines
for rn, ko in got.items():
    if rn not in targets:
        continue
    jp = targets[rn]
    ko = ko.replace('_x000D_', '').strip()
    ko = ELL.sub(r'\1 ', ko)
    if not ko:
        fails['빈값'].append(rn); continue
    bad = False
    for c in CIRC:
        if jp.count(c) != ko.count(c):
            fails['변수불일치'].append(rn); bad = True; break
    if bad: continue
    if JPCH.search(ko):
        fails['일본어잔존'].append(rn); continue
    oux = [ch for ch in ko if '가' <= ch <= '힣' and ch not in KS]
    if oux:
        fails['폰트밖글자'].append((rn, ''.join(sorted(set(oux))))); continue
    # 원문 자체가 '화자「대사」' 형태면 번역도 그래야 정상. 원문에 없을 때만 결함.
    if re.match(r'^[가-힣A-Za-z]{1,6}「', ko) and '「' not in jp:
        fails['화자프리픽스'].append(rn); continue
    jpl = jp.count('\n') + 1
    lines, _ = fit_turn_lines(ko, jpl, 7, speaker='')
    if len(lines) > jpl + 3:
        warn_len.append((rn, jpl, len(lines)))
    ok[rn] = ko

print(f"\n검증 통과 {len(ok)} / 실패 {sum(len(v) for v in fails.values())}")
for k, v in fails.items():
    print(f"  {k}: {len(v)}  예: {v[:6]}")
print(f"길이 경고(>원문+3줄): {len(warn_len)}")

# 재시도 목록 저장
retry = sorted(set(missing) | {x if isinstance(x, int) else x[0]
                               for v in fails.values() for x in v})
json.dump(retry, open(os.path.join(DIR, '_retry_rns.json'), 'w'), indent=0)
print(f"재시도 필요: {len(retry)}행 → _retry_rns.json")

if WRITE:
    n = 0
    for rn, ko in ok.items():
        ws.cell(rn, 10).value = ko
        n += 1
    wb.save(os.path.join(HERE, 'srwj_matched_all_0625.xlsx'))
    print(f"✓ 엑셀 반영 {n}행")
else:
    print("(검증 전용 — 반영은 --write)")
