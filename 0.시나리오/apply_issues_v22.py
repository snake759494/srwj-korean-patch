# -*- coding: utf-8 -*-
"""GitHub 이슈 #5·#8·#21~#29 반영 — v2.2

세 데이터 파일(시나리오 xlsx J열 / battle_dialogue.json ko / translations.json ko)에
제보 내용을 한 번에 적용한다. 항목마다 근거 이슈 번호를 달아 두었다.

  python apply_issues_v22.py            # 미리보기(변경 안 함)
  python apply_issues_v22.py --write    # 실제 반영

원칙
  · 일본어 원문(jp)이 구분하는 표기는 원문을 따른다(예: ニュートロンジャマー vs Ｎジャマー).
  · 폭 규칙(전투 14칸·3줄, 시나리오 자동 줄바꿈)을 넘기지 않는 범위에서만 고친다.
  · 연속 공백은 삽입기가 전각 공백 2칸으로 그려 화면에 큰 빈칸이 생기므로 한 칸으로 줄인다(#8).
"""
import json, os, re, sys

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(HERE)
XLSX = os.path.join(HERE, 'srwj_matched_all_0625.xlsx')
BATTLE = os.path.join(ROOT, '2. 전투대사패치', 'battle_dialogue.json')
TRANS = os.path.join(ROOT, '3. SJIS추출', 'translations.json')

WRITE = '--write' in sys.argv
log = []


def note(kind, where, before, after, iss):
    log.append((kind, where, before, after, iss))


# ══════════════════════════════════════════════════════════════════
# 1. 시나리오 (xlsx J열)
# ══════════════════════════════════════════════════════════════════
# (a) 행 지정 치환 — 제보에서 문장이 특정된 것
XLSX_EXACT = [
    # (행, 옛문구, 새문구, 이슈)
    (19193, '카펜타리아', '카펜테리아', '#5'),
    (7153, '혼푼', '본푼', '#22'),
    (7170, '혼푼', '본푼', '#22'),
    (28500, '없앨지도', '죽일지도', '#24'),
    (28683, '개인적인 호불호는', '개인적인 악감정은', '#24'),
    (29288, '이게 마지막이에요', '이게 마지막입니다', '#25'),
    (28481, '타전해라', '연락해라', '#26'),
    (30866, '타전해', '연락해', '#26'),
    (31639, '그 자식 또한', '그의 자식 또한', '#27'),
    (33165, '고, 괜찮아', '괘, 괜찮아', '#29'),
    (33091, '잔적 수 제로', '남은 적 없음', '#29'),
    (33107, '쿄우지 형님', '쿄우지 형', '#29'),
    (33303, '채로는아아앗', '채로느으으은', '#29'),
    (33402, '저 처녀를', '저 아가씨를', '#29'),
    (33671, '분명히 알았다', '분명히 들었다', '#29'),
    (32967, '급속히 증대', '급속 상승', '#29'),
]

# (b) 전역 치환 — 용어 통일
XLSX_GLOBAL = [
    ('바지룰', '버지룰', '#24'),          # 나탈 버지룰(Natarle Badgiruel)
    ('빅 팰컨', '빅 팔콘', '#27'),        # 볼테스V 기지명 통일
    ('어머니 되는 함선', '모선', '#29'),
    ('한때의 체면', '순간의 체면', '#29'),
    ('짓임을 알라!', '짓임을 알아라!', '#29'),
    ('행동임을 알라!', '행동임을 알아라!', '#29'),
    ('기사도 불각오', '기사도에 어긋난다', '#29'),
    ('가 다오.', '가거라.', '#29'),
]

# (c) 원문이 Ｎジャマー 인 곳만 'N재머'로 (원문이 ニュートロンジャマー면 그대로) — #21 #24 #27
NJAMMER_JP = 'Ｎジャマー'


def patch_xlsx():
    from openpyxl import load_workbook
    wb = load_workbook(XLSX)
    ws = wb[wb.sheetnames[0]]
    JCOL, KCOL = 10, 11          # J=한국어, K=일본어(ROM)
    changed = 0

    for row, old, new, iss in XLSX_EXACT:
        c = ws.cell(row=row, column=JCOL)
        v = c.value
        if v and old in str(v):
            nv = str(v).replace(old, new)
            note('시나리오', 'R%d' % row, str(v), nv, iss)
            if WRITE:
                c.value = nv
            changed += 1
        else:
            note('시나리오', 'R%d' % row, '(못 찾음: %s)' % old, '-', iss)

    for r in range(2, ws.max_row + 1):
        c = ws.cell(row=r, column=JCOL)
        v = c.value
        if not v:
            continue
        s = nv = str(v)
        tags = []
        for old, new, iss in XLSX_GLOBAL:
            if old in nv:
                nv = nv.replace(old, new)
                tags.append(iss)
        # Ｎジャマー 원문 대응
        jp = ws.cell(row=r, column=KCOL).value or ''
        if NJAMMER_JP in str(jp) and '뉴트론 재머' in nv:
            nv = nv.replace('뉴트론 재머', 'N재머')
            tags.append('#21')
        # 연속 공백 정리 (#8)
        nv2 = re.sub(r'[ 　]{2,}', ' ', nv)
        if nv2 != nv:
            nv = nv2
            tags.append('#8')
        if nv != s:
            note('시나리오', 'R%d' % r, s, nv, ','.join(sorted(set(tags))))
            if WRITE:
                c.value = nv
            changed += 1

    if WRITE:
        wb.save(XLSX)
    return changed


# ══════════════════════════════════════════════════════════════════
# 2. 전투 대사 (battle_dialogue.json ko)
# ══════════════════════════════════════════════════════════════════
BATTLE_EXACT = [
    # (인덱스, 새 ko, 이슈, 설명)
    (2812, '컨, 컨티뉴만\n할 수 있으면…!', '#25', 'JP コ、コンティニュー — 첫 음절 일치'),
    (5399, '하이코트!\n볼테카아아앗!!', '#26', '하이코트오 → 하이코트'),
    (5401, '하이코트!!\n볼테카아아앗!!', '#26', '하이코트오오 → 하이코트'),
    (9914, '료! 접근전은\n너에게 맡기겠어!!', '#26', '뒤 대사와 이어지도록'),
    (1961, '여기는 아스란\n자라다, 공격을\n개시한다!', '#27', '무선 호출 관용 표현'),
    (3124, '이걸로\n마지막이다…!', '#27', '제오라이머 명왕공격'),
    (9911, '해치워주마!', '#27', '단쿠가 やってやるぜ'),
    (10262, '그 틈은\n놓치지 않는다!\n해치워주마!!', '#27', '단쿠가 やってやるぜ'),
    (3745, '여,\n여기까지인가!?', '#27', 'JP こ、ここまでか — 말더듬 첫 음절'),
    (7252, '여,\n여기까지인가…', '#27', '동일'),
    (10359, '여,\n여기까지인가!!', '#29', '동일'),
    (11637, '절망의 불꽃으로\n타 죽어라! 공포의\n어둠 속에서 얼어붙어라!', '#27', '띄어쓰기 복원'),
    (368, '필-살!', '#29', 'JP ひいっさつ'),
    (11656, '으오오오!\n흥! 이 몸이이이아아아!', '#29', 'JP 何のぉおお'),
    (11664, '안\n통한다아아앗!', '#29', 'JP 効ぃかぁぬぅうぅわぁあぁ'),
    (1521, '사선상에\n아군기 없음!', '#29', '사선축 → 사선상(다른 대사와 통일)'),
    (31, '붉게\n타오르고 있다!', '#29', 'JP 真っ赤に燃える'),
    (49, '붉게\n타오르고 있다!', '#29', '동일'),
]

BATTLE_GLOBAL = [
    ('빅 팰컨', '빅 팔콘', '#27'),
]


def patch_battle():
    d = json.load(open(BATTLE, encoding='utf-8'))
    ent = d['entries']
    changed = 0
    for idx, new, iss, why in BATTLE_EXACT:
        old = ent[idx].get('ko') or ''
        if old == new:
            continue
        note('전투', 'i%d blk%s' % (idx, ent[idx]['blk']), old, new, '%s (%s)' % (iss, why))
        if WRITE:
            ent[idx]['ko'] = new
        changed += 1
    for i, e in enumerate(ent):
        ko = e.get('ko') or ''
        if not ko:
            continue
        nv = ko
        tags = []
        for old, new, iss in BATTLE_GLOBAL:
            if old in nv:
                nv = nv.replace(old, new)
                tags.append(iss)
        nv2 = re.sub(r'[ 　]{2,}', ' ', nv)     # #8 빈칸 과다
        if nv2 != nv:
            nv = nv2
            tags.append('#8')
        if nv != ko and not any(i == x[0] for x in BATTLE_EXACT):
            note('전투', 'i%d blk%s' % (i, e['blk']), ko, nv, ','.join(sorted(set(tags))))
            if WRITE:
                e['ko'] = nv
            changed += 1
    if WRITE:
        json.dump(d, open(BATTLE, 'w', encoding='utf-8'), ensure_ascii=False, indent=1)
    return changed


# ══════════════════════════════════════════════════════════════════
# 3. 시스템 텍스트 (translations.json ko)
# ══════════════════════════════════════════════════════════════════
# 무장명 단위 통일 — 이미 ｍｍ/Ｍ 로 바뀐 것들과 표기를 맞춘다 (#20 후속, #25)
UNIT_RULES = [
    (re.compile(r'(?<=[０-９0-9])(밀리|미리)'), 'ｍｍ', 'ミリ'),
    (re.compile(r'(?<=[０-９0-9])센티'), 'ｃｍ', 'センチ'),
    (re.compile(r'(?<=[０-９0-9])미터'), 'Ｍ', 'メートル'),
]
TRANS_EXACT = [
    ('달상공', '달 상공', '#27'),
    ('거주구', '거주구역', '#27'),
]
TRANS_GLOBAL = [('빅 팰컨', '빅 팔콘', '#27')]


def enc_len(s):
    """SJIS 기준 바이트 길이(한글·전각 2바이트, 반각 1바이트)."""
    n = 0
    for ch in s:
        n += 1 if (ord(ch) < 0x80) else 2
    return n


def patch_trans():
    d = json.load(open(TRANS, encoding='utf-8'))
    ent = d['entries']
    changed = 0
    for e in ent:
        ko, jp = e.get('ko') or '', e.get('jp') or ''
        if not ko:
            continue
        nv = ko
        tags = []
        for rx, rep, jpmark in UNIT_RULES:
            if jpmark in jp and rx.search(nv):
                nv = rx.sub(rep, nv)
                tags.append('#25')
        for old, new, iss in TRANS_GLOBAL:
            if old in nv:
                nv = nv.replace(old, new)
                tags.append(iss)
        for old, new, iss in TRANS_EXACT:
            if nv == old:
                nv = new
                tags.append(iss)
        if nv != ko:
            note('시스템', '%s len=%s→%s' % (e['off'], enc_len(ko), enc_len(nv)), ko, nv,
                 ','.join(sorted(set(tags))))
            if WRITE:
                e['ko'] = nv
            changed += 1
    if WRITE:
        json.dump(d, open(TRANS, 'w', encoding='utf-8'), ensure_ascii=False, indent=1)
    return changed


if __name__ == '__main__':
    sys.stdout.reconfigure(encoding='utf-8')
    print('모드:', '실제 반영(--write)' if WRITE else '미리보기')
    n1 = patch_xlsx()
    n2 = patch_battle()
    n3 = patch_trans()
    by = {}
    for kind, where, before, after, iss in log:
        by.setdefault(kind, []).append((where, before, after, iss))
    for kind in ('시나리오', '전투', '시스템'):
        rows = by.get(kind, [])
        print('\n' + '=' * 78)
        print('%s — %d건' % (kind, len(rows)))
        print('=' * 78)
        for where, before, after, iss in rows[:400]:
            print('  [%s] %s' % (iss, where))
            print('     - %s' % before.replace('\n', '⏎'))
            print('     + %s' % after.replace('\n', '⏎'))
        if len(rows) > 400:
            print('  ... 외 %d건' % (len(rows) - 400))
    print('\n합계: 시나리오 %d / 전투 %d / 시스템 %d' % (n1, n2, n3))
