# -*- coding: utf-8 -*-
"""이슈 #1(10화) 본문 수정 — v2.4

줄맞춤은 줄바꿈 규칙 자체를 고쳐 해결하고(srwj_wrap.balanced_wrap,
2. 전투대사패치/rewrap_balance.py), 여기서는 글을 바꿔야 하는 항목만 다룬다.

  python apply_issues_v24.py            # 미리보기
  python apply_issues_v24.py --write    # 반영
"""
import json, os, sys

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(HERE)
XLSX = os.path.join(HERE, 'srwj_matched_all_0625.xlsx')
BATTLE = os.path.join(ROOT, '2. 전투대사패치', 'battle_dialogue.json')
WRITE = '--write' in sys.argv
log = []

# (행, 옛문구, 새문구, 항목번호, 근거)
XLSX_FIX = [
    (3864, '문제는 없는 거 같네\n도와줘서 고마워요',
     '문제는 없는 거 같네, 도와줘서 고마워요', 5, '두 문장을 한 대사로 (원문 問題はなさそうね。助かります)'),
    (3865, '뭘 그런거 가지고\n', '뭘 그런거 가지고,\n', 6, '원문 なぁに、 의 쉼표'),
    (3907, '라담 라담 아주 노래를 부르네\n라담오타쿠 진짜 극혐이구만',
     '라담 라담 아주 노래를 부르는군\n라담 오타쿠는 질색이야', 10,
     '원문 嫌だねぇラダムオタクは — 유행어(극혐) 대신 사전적 표현'),
    (3972, '... 스테이시스(정지) 확인', '... 스테이시스 확인', 12,
     '원문 ・・・ステイシス確認。 — 덧붙인 풀이말 제거(다른 대사도 스테이시스 그대로 씀)'),
    (4149, '저게 ⑤ 만들고', '저게 ⑤를 만들고', 15, '원문 あれが⑤を作って — 목적격 조사 누락'),
    (4152, '아무래도 성가신 적인 모양이군, 저건.', '저건 아무래도 성가신 적인 모양이군.', 16,
     '원문 どうもやっかいな敵らしいなあれは — 어순 정리'),
    (4176, '야 잠깐만... ', '이봐, 잠깐만... ', 18, '원문 おい待てよ'),
    (4344, '목성 도마뱀에게는 애를 먹고 있어', '목성 도마뱀에게는 고전하고 있어', 43,
     '원문 手を焼いている'),
]

# (인덱스, 옛 ko, 새 ko, 항목번호, 근거)
BATTLE_FIX = [
    (1231, '사정거리 안에\n있다니!', '사정거리 안에\n들어오다니!', 20,
     '원문 射程内にいるとはなっ！ — 무장 표기는 사정거리로 통일'),
    (10852, '두지\n않겠다고… 말했잖아!!', '가만두지\n않겠다고… 말했잖아!!', 21,
     '원문 やらせないって・・・いっただろ！！ — 목적어 없이 끊겨 있었음'),
]


def patch_xlsx():
    from openpyxl import load_workbook
    wb = load_workbook(XLSX)
    ws = wb[wb.sheetnames[0]]
    n = 0
    for row, old, new, no, why in XLSX_FIX:
        c = ws.cell(row=row, column=10)
        v = str(c.value or '')
        if old in v:
            nv = v.replace(old, new)
            log.append(('시나리오', 'R%d' % row, v, nv, no, why))
            if WRITE:
                c.value = nv
            n += 1
        else:
            log.append(('시나리오', 'R%d' % row, '(못 찾음) %s' % old, '-', no, why))
    if WRITE:
        wb.save(XLSX)
    return n


def patch_battle():
    d = json.load(open(BATTLE, encoding='utf-8'))
    ent = d['entries']
    n = 0
    for idx, old, new, no, why in BATTLE_FIX:
        cur = ent[idx].get('ko') or ''
        if cur == old:
            log.append(('전투', 'i%d blk%s' % (idx, ent[idx]['blk']), cur, new, no, why))
            if WRITE:
                ent[idx]['ko'] = new
            n += 1
        else:
            log.append(('전투', 'i%d' % idx, '(현재값 다름) %s' % cur.replace('\n', ' ⏎ '), new, no, why))
    if WRITE:
        json.dump(d, open(BATTLE, 'w', encoding='utf-8'), ensure_ascii=False, indent=1)
    return n


if __name__ == '__main__':
    sys.stdout.reconfigure(encoding='utf-8')
    print('모드:', '반영(--write)' if WRITE else '미리보기')
    a = patch_xlsx()
    b = patch_battle()
    for kind, where, before, after, no, why in log:
        print('\n[%2d] %s %s — %s' % (no, kind, where, why))
        print('   - %s' % before.replace('\n', ' ⏎ '))
        print('   + %s' % after.replace('\n', ' ⏎ '))
    print('\n합계: 시나리오 %d / 전투 %d' % (a, b))
