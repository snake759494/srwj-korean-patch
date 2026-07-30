# -*- coding: utf-8 -*-
"""AI 번역 행(I열 빈) 재번역용 청크 생성.

각 청크 = 연속 턴 구간 전체(문맥) + 그중 재번역 대상 행 표시 + 관련 용어집.
출력: scratchpad/retrans/chunk_NNN.json + index.json
"""
import sys, os, json, re, collections

sys.stdout.reconfigure(encoding='utf-8')
HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(HERE)
OUT = sys.argv[1]
MAX_TGT, MAX_SPAN, CTX = 32, 110, 6

from openpyxl import load_workbook
ws = load_workbook(os.path.join(HERE, 'srwj_matched_all_0625.xlsx'),
                   read_only=True, data_only=True)['매칭 결과']
rows = list(ws.iter_rows(min_row=2, values_only=True))

# 용어집: translations.json 에서 (jp→ko), 2~14자·제어문자 없는 항목
tr = json.load(open(os.path.join(ROOT, '3. SJIS추출', 'translations.json'),
                    encoding='utf-8'))['entries']
TERM = {}
for e in tr:
    jp, ko = e['jp'], e['ko']
    if (2 <= len(jp) <= 14 and ko and not re.search(r'[#%&<>|*\d]', jp)
            and not re.search(r'[#%&<>|]', ko)
            and re.search(r'[ァ-ヺー一-鿿]', jp)):
        TERM.setdefault(jp, ko)
print(f"용어집 후보: {len(TERM)}")

# 아카이브별 턴 목록
by_arc = collections.defaultdict(list)
for i, r in enumerate(rows, 2):
    if r[0] is None or r[1] is None:
        continue
    ko = str(r[9]) if r[9] else ''
    jp = str(r[10]) if r[10] else ''
    if not ko.strip() or not jp.strip():
        continue
    human = bool(r[8] and str(r[8]).strip())
    by_arc[int(r[0])].append(dict(rn=i, t=int(r[1]), cid=str(r[4] or ''),
                                  spk=str(r[6] or ''), jp=jp, ko=ko,
                                  human=human))

os.makedirs(OUT, exist_ok=True)
chunks = []
for arc in sorted(by_arc):
    turns = by_arc[arc]
    n = len(turns)
    i = 0
    while i < n:
        # 다음 대상 찾기
        while i < n and turns[i]['human']:
            i += 1
        if i >= n:
            break
        start = max(0, i - CTX)
        j = i
        tgt = 0
        while j < n and (j - start) < MAX_SPAN and tgt < MAX_TGT:
            if not turns[j]['human']:
                tgt += 1
            j += 1
        end = min(n, j + 4)              # 뒤 문맥
        seg = turns[start:end]
        targets = [t['rn'] for t in seg[i - start: j - start] if not t['human']]
        chunks.append(dict(arc=arc, turns=seg, targets=targets))
        i = j

# 청크별 용어집 축소
for ci, ch in enumerate(chunks):
    text = ''.join(t['jp'] for t in ch['turns'])
    g = {jp: ko for jp, ko in TERM.items() if jp in text}
    ch['glossary'] = g
    json.dump(ch, open(os.path.join(OUT, f'chunk_{ci:03d}.json'), 'w',
                       encoding='utf-8'), ensure_ascii=False, indent=1)

tot = sum(len(c['targets']) for c in chunks)
json.dump({'n_chunks': len(chunks), 'total_targets': tot,
           'chunks': [{'i': i, 'arc': c['arc'], 'targets': len(c['targets'])}
                      for i, c in enumerate(chunks)]},
          open(os.path.join(OUT, 'index.json'), 'w', encoding='utf-8'),
          ensure_ascii=False, indent=1)
print(f"청크 {len(chunks)}개, 대상 {tot}행 → {OUT}")
